from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor, Twips


ROOT = Path(r"D:\_HQZOZONE")
OUTPUT = ROOT / "reports" / "csg_group_meeting_report_2026-08-26.docx"
IMAGE_OVERALL = Path(
    r"C:\Users\29737\AppData\Local\Temp\codex-clipboard-4cb30153-cdde-48a5-b61c-bafc71fe033b.png"
)
IMAGE_STEPS = Path(
    r"C:\Users\29737\AppData\Local\Temp\codex-clipboard-b588e1a7-ff38-42b1-82e7-905d2c72eccb.png"
)
IMAGE_ARCHITECTURE = Path(
    r"C:\Users\29737\AppData\Local\Temp\codex-clipboard-ec8e412a-73d6-4739-8013-471ae85ff815.png"
)

FONT_BODY = "Microsoft YaHei"
FONT_LATIN = "Calibri"
INK = "1F2937"
MUTED = "667085"
BLUE = "1F77B4"
BLUE_DARK = "1F4D78"
BLUE_PALE = "EAF3FB"
ORANGE = "FF7F0E"
GREEN = "2CA02C"
GREEN_PALE = "EAF6EC"
GRAY_PALE = "F2F4F7"
GRAY_RULE = "D8DEE6"
AMBER_PALE = "FFF5E8"
WHITE = "FFFFFF"

# A4 report override to the selected standard_business_brief preset.
CONTENT_WIDTH_DXA = 9864


def rgb(hex_color):
    return RGBColor.from_string(hex_color)


def set_run_font(run, size=None, color=INK, bold=None, italic=None, latin=FONT_LATIN):
    run.font.name = latin
    run.font.color.rgb = rgb(color)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.rFonts
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.insert(0, r_fonts)
    r_fonts.set(qn("w:ascii"), latin)
    r_fonts.set(qn("w:hAnsi"), latin)
    r_fonts.set(qn("w:eastAsia"), FONT_BODY)
    lang = r_pr.find(qn("w:lang"))
    if lang is None:
        lang = OxmlElement("w:lang")
        r_pr.append(lang)
    lang.set(qn("w:eastAsia"), "zh-CN")


def set_style_font(style, size, color=INK, bold=None):
    style.font.name = FONT_LATIN
    style.font.size = Pt(size)
    style.font.color.rgb = rgb(color)
    if bold is not None:
        style.font.bold = bold
    style.font.italic = False
    style.font.underline = False
    r_pr = style.element.get_or_add_rPr()
    r_fonts = r_pr.rFonts
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.insert(0, r_fonts)
    r_fonts.set(qn("w:ascii"), FONT_LATIN)
    r_fonts.set(qn("w:hAnsi"), FONT_LATIN)
    r_fonts.set(qn("w:eastAsia"), FONT_BODY)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in {
        "top": top,
        "start": start,
        "bottom": bottom,
        "end": end,
    }.items():
        node = tc_mar.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_table_geometry(table, widths_dxa, indent_dxa=120):
    if sum(widths_dxa) != CONTENT_WIDTH_DXA:
        raise ValueError(f"table widths must sum to {CONTENT_WIDTH_DXA}")

    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    tbl_w.set(qn("w:type"), "dxa")
    tbl_w.set(qn("w:w"), str(CONTENT_WIDTH_DXA))

    tbl_ind = tbl_pr.first_child_found_in("w:tblInd")
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:type"), "dxa")
    tbl_ind.set(qn("w:w"), str(indent_dxa))

    layout = tbl_pr.first_child_found_in("w:tblLayout")
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        prevent_row_split(row)
        for index, cell in enumerate(row.cells):
            cell.width = Twips(widths_dxa[index])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            tc_w = cell._tc.get_or_add_tcPr().get_or_add_tcW()
            tc_w.set(qn("w:type"), "dxa")
            tc_w.set(qn("w:w"), str(widths_dxa[index]))


def set_table_borders(table, color=GRAY_RULE, size="6"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def style_table(table, widths_dxa, header_fill=GRAY_PALE):
    set_table_geometry(table, widths_dxa)
    set_table_borders(table)
    set_repeat_table_header(table.rows[0])
    for cell in table.rows[0].cells:
        set_cell_shading(cell, header_fill)
        for paragraph in cell.paragraphs:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(0)
            for run in paragraph.runs:
                set_run_font(run, size=9.5, bold=True)
    for row in table.rows[1:]:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 1.05
                for run in paragraph.runs:
                    set_run_font(run, size=9.2)


def shade_paragraph(paragraph, fill, left_border=None):
    p_pr = paragraph._p.get_or_add_pPr()
    shading = p_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        p_pr.append(shading)
    shading.set(qn("w:fill"), fill)
    if left_border:
        p_bdr = p_pr.find(qn("w:pBdr"))
        if p_bdr is None:
            p_bdr = OxmlElement("w:pBdr")
            p_pr.append(p_bdr)
        left = OxmlElement("w:left")
        left.set(qn("w:val"), "single")
        left.set(qn("w:sz"), "18")
        left.set(qn("w:space"), "8")
        left.set(qn("w:color"), left_border)
        p_bdr.append(left)


def add_body(document, text, bold_prefix=None, after=6, keep=False):
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = 1.10
    paragraph.paragraph_format.keep_with_next = keep
    if bold_prefix and text.startswith(bold_prefix):
        lead = paragraph.add_run(bold_prefix)
        set_run_font(lead, size=11, bold=True)
        tail = paragraph.add_run(text[len(bold_prefix) :])
        set_run_font(tail, size=11)
    else:
        run = paragraph.add_run(text)
        set_run_font(run, size=11)
    return paragraph


def add_bullet(document, text, accent=None, after=5):
    paragraph = document.add_paragraph(style="List Bullet")
    paragraph.paragraph_format.left_indent = Inches(0.5)
    paragraph.paragraph_format.first_line_indent = Inches(-0.25)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = 1.167
    run = paragraph.add_run(text)
    set_run_font(run, size=10.6, color=accent or INK)
    return paragraph


def add_numbered(document, title, body, number):
    paragraph = document.add_paragraph(style="List Number")
    paragraph.paragraph_format.left_indent = Inches(0.5)
    paragraph.paragraph_format.first_line_indent = Inches(-0.25)
    paragraph.paragraph_format.space_after = Pt(7)
    paragraph.paragraph_format.line_spacing = 1.167
    run = paragraph.add_run(f"{title}：")
    set_run_font(run, size=10.6, bold=True, color=BLUE_DARK)
    run = paragraph.add_run(body)
    set_run_font(run, size=10.6)
    return paragraph


def add_heading(document, number, title, level=1):
    paragraph = document.add_paragraph(style=f"Heading {level}")
    run = paragraph.add_run(f"{number}  {title}" if number else title)
    set_run_font(
        run,
        size=16 if level == 1 else 13 if level == 2 else 12,
        color=BLUE if level < 3 else BLUE_DARK,
        bold=True,
    )
    paragraph.paragraph_format.keep_with_next = True
    return paragraph


def add_kicker(document, text, color=ORANGE, align=WD_ALIGN_PARAGRAPH.LEFT, after=6):
    paragraph = document.add_paragraph(style="Kicker")
    paragraph.alignment = align
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(after)
    run = paragraph.add_run(text.upper())
    set_run_font(run, size=9.5, color=color, bold=True)
    return paragraph


def add_callout(document, label, body, accent=BLUE, fill=BLUE_PALE):
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.left_indent = Cm(0.25)
    paragraph.paragraph_format.right_indent = Cm(0.15)
    paragraph.paragraph_format.space_before = Pt(5)
    paragraph.paragraph_format.space_after = Pt(8)
    paragraph.paragraph_format.line_spacing = 1.12
    shade_paragraph(paragraph, fill, accent)
    label_run = paragraph.add_run(f"  {label}  ")
    set_run_font(label_run, size=10.3, color=accent, bold=True)
    body_run = paragraph.add_run(body)
    set_run_font(body_run, size=10.3)
    return paragraph


def add_equation_box(document, lines):
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.left_indent = Cm(0.55)
    paragraph.paragraph_format.right_indent = Cm(0.55)
    paragraph.paragraph_format.space_before = Pt(7)
    paragraph.paragraph_format.space_after = Pt(9)
    paragraph.paragraph_format.line_spacing = 1.35
    shade_paragraph(paragraph, "F7F9FC", BLUE)
    for index, line in enumerate(lines):
        run = paragraph.add_run(("\n" if index else "") + line)
        set_run_font(run, size=11.2, color=BLUE_DARK, bold=True, latin="Consolas")
    return paragraph


def set_alt_text(inline_shape, title, description):
    doc_pr = inline_shape._inline.docPr
    doc_pr.set("title", title)
    doc_pr.set("descr", description)


def add_figure(document, path, caption, alt_text, width_cm):
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(5)
    paragraph.paragraph_format.space_after = Pt(2)
    paragraph.paragraph_format.keep_with_next = True
    shape = paragraph.add_run().add_picture(str(path), width=Cm(width_cm))
    set_alt_text(shape, caption, alt_text)

    caption_paragraph = document.add_paragraph(style="Caption")
    caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption_paragraph.paragraph_format.space_before = Pt(2)
    caption_paragraph.paragraph_format.space_after = Pt(1)
    caption_paragraph.paragraph_format.keep_with_next = True
    caption_run = caption_paragraph.add_run(caption)
    set_run_font(caption_run, size=9, color=INK, bold=True)

    source = document.add_paragraph()
    source.alignment = WD_ALIGN_PARAGRAPH.CENTER
    source.paragraph_format.space_before = Pt(0)
    source.paragraph_format.space_after = Pt(6)
    source_run = source.add_run("来源：本次实验平台截图。")
    set_run_font(source_run, size=8.2, color=MUTED)
    return shape


def add_page_break(document):
    document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def add_page_field(run):
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend((begin, instruction, separate, text, end))


def configure_styles(document):
    normal = document.styles["Normal"]
    set_style_font(normal, 11, INK)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    title = document.styles["Title"]
    set_style_font(title, 30, BLUE_DARK, True)
    title.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(8)
    title_p_pr = title.element.get_or_add_pPr()
    title_border = title_p_pr.find(qn("w:pBdr"))
    if title_border is not None:
        title_p_pr.remove(title_border)

    subtitle = document.styles["Subtitle"]
    set_style_font(subtitle, 14, MUTED, True)
    subtitle.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_before = Pt(0)
    subtitle.paragraph_format.space_after = Pt(24)

    kicker = document.styles.add_style("Kicker", WD_STYLE_TYPE.PARAGRAPH)
    set_style_font(kicker, 9.5, ORANGE, True)
    kicker.paragraph_format.space_before = Pt(0)
    kicker.paragraph_format.space_after = Pt(6)

    h1 = document.styles["Heading 1"]
    set_style_font(h1, 16, BLUE, True)
    h1.paragraph_format.space_before = Pt(16)
    h1.paragraph_format.space_after = Pt(8)
    h1.paragraph_format.keep_with_next = True

    h2 = document.styles["Heading 2"]
    set_style_font(h2, 13, BLUE, True)
    h2.paragraph_format.space_before = Pt(12)
    h2.paragraph_format.space_after = Pt(6)
    h2.paragraph_format.keep_with_next = True

    h3 = document.styles["Heading 3"]
    set_style_font(h3, 12, BLUE_DARK, True)
    h3.paragraph_format.space_before = Pt(8)
    h3.paragraph_format.space_after = Pt(4)
    h3.paragraph_format.keep_with_next = True

    caption = document.styles["Caption"]
    set_style_font(caption, 9, INK)
    caption.paragraph_format.space_before = Pt(2)
    caption.paragraph_format.space_after = Pt(1)

    for style_name in ("List Bullet", "List Number"):
        style = document.styles[style_name]
        set_style_font(style, 10.6, INK)
        style.paragraph_format.left_indent = Inches(0.5)
        style.paragraph_format.first_line_indent = Inches(-0.25)
        style.paragraph_format.space_after = Pt(8)
        style.paragraph_format.line_spacing = 1.167


def configure_page(document):
    section = document.sections[0]
    section.start_type = WD_SECTION.NEW_PAGE
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(1.8)
    section.right_margin = Cm(1.8)
    section.top_margin = Cm(1.7)
    section.bottom_margin = Cm(1.7)
    section.header_distance = Cm(0.65)
    section.footer_distance = Cm(0.65)

    header = section.header
    header_paragraph = header.paragraphs[0]
    header_paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header_paragraph.paragraph_format.space_after = Pt(0)
    header_run = header_paragraph.add_run("CSG | 火星臭氧多步预测组会汇报")
    set_run_font(header_run, size=8.2, color=MUTED, bold=True)

    footer = section.footer
    footer_paragraph = footer.paragraphs[0]
    footer_paragraph.paragraph_format.tab_stops.add_tab_stop(
        Cm(17.35), WD_TAB_ALIGNMENT.RIGHT
    )
    footer_paragraph.paragraph_format.space_before = Pt(0)
    footer_paragraph.paragraph_format.space_after = Pt(0)
    left_run = footer_paragraph.add_run("2026-08-26  ·  INTERNAL RESEARCH BRIEF")
    set_run_font(left_run, size=8, color=MUTED)
    footer_paragraph.add_run("\t")
    page_run = footer_paragraph.add_run()
    set_run_font(page_run, size=8, color=MUTED, bold=True)
    add_page_field(page_run)


def add_metric_strip(document, metrics):
    table = document.add_table(rows=1, cols=len(metrics))
    widths = [CONTENT_WIDTH_DXA // len(metrics)] * len(metrics)
    widths[-1] += CONTENT_WIDTH_DXA - sum(widths)
    set_table_geometry(table, widths, indent_dxa=120)
    set_table_borders(table, color=WHITE, size="0")
    set_repeat_table_header(table.rows[0])
    for index, (value, label, color) in enumerate(metrics):
        cell = table.cell(0, index)
        set_cell_shading(cell, "F7F9FC")
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_before = Pt(2)
        paragraph.paragraph_format.space_after = Pt(2)
        value_run = paragraph.add_run(value)
        set_run_font(value_run, size=17, color=color, bold=True)
        label_run = paragraph.add_run(f"\n{label}")
        set_run_font(label_run, size=8.5, color=MUTED, bold=True)
    spacer = document.add_paragraph()
    spacer.paragraph_format.space_after = Pt(0)
    return table


def add_cover(document):
    for _ in range(4):
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(12)

    add_kicker(
        document,
        "GROUP MEETING · MODEL REVIEW",
        color=ORANGE,
        align=WD_ALIGN_PARAGRAPH.CENTER,
        after=16,
    )
    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(8)
    run = title.add_run("CSG 条件时空门控模型")
    set_run_font(run, size=30, color=BLUE_DARK, bold=True)

    subtitle = document.add_paragraph(style="Subtitle")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(24)
    run = subtitle.add_run("融合太阳黄经与 MOLA 地形的火星臭氧多步预测")
    set_run_font(run, size=14, color=MUTED, bold=True)

    add_callout(
        document,
        "本次汇报",
        "解释 CSG 的结构、联合门控机制与实验表现，并明确当前证据能够支持的结论以及下一轮实验重点。",
        accent=BLUE,
        fill=BLUE_PALE,
    )

    for _ in range(2):
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(10)

    add_metric_strip(
        document,
        [
            ("20 → 20", "历史窗口 / 预测步", BLUE),
            ("+8.19%", "相对基线参数开销", ORANGE),
            ("3 条件源", "动态特征 / Ls / 地形", GREEN),
        ],
    )

    date = document.add_paragraph()
    date.alignment = WD_ALIGN_PARAGRAPH.CENTER
    date.paragraph_format.space_before = Pt(28)
    run = date.add_run("组会汇报文档  |  2026 年 8 月 26 日")
    set_run_font(run, size=10, color=MUTED, bold=True)


def add_research_page(document):
    add_kicker(document, "RESEARCH QUESTION")
    add_heading(document, "01", "研究问题与核心结论")
    add_callout(
        document,
        "核心问题",
        "在保持 ConvLSTM-SimVP 主干不变的前提下，引入历史太阳黄经 Ls 与静态 MOLA 地形，能否降低火星臭氧多步预测误差？",
        accent=BLUE,
        fill=BLUE_PALE,
    )

    add_body(
        document,
        "CSG 将时变的大气动力特征、季节相位和静态地形先验放入同一个轻量门控模块中。这样的设计把实验边界控制得较干净：主干编码器、时间翻译器和解码器与基线保持一致，性能差异主要来自条件门控。",
    )

    add_heading(document, "", "对比对象", level=2)
    table = document.add_table(rows=1, cols=3)
    for idx, text in enumerate(("模型", "输入与结构", "对比作用")):
        table.rows[0].cells[idx].text = text
    rows = [
        ("SimVP", "纯视频预测式时空建模", "基础时空预测基线"),
        ("ConvLSTM-SimVP", "ConvLSTM 历史编码 + 轻量 SimVP", "检验循环时序编码的增益"),
        ("CSG", "主干不变 + Ls/MOLA 联合条件门", "检验外部物理条件的增益"),
    ]
    for values in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(values):
            cells[idx].text = value
    style_table(table, [2000, 4300, 3564])

    add_heading(document, "", "当前结果传递的三个信号", level=2)
    add_bullet(document, "总体 RMSE 排名为 CSG < ConvLSTM-SimVP < SimVP，方向与研究假设一致。")
    add_bullet(document, "CSG 的优势贯穿多数预测步，长预测步的分离更明显。")
    add_bullet(document, "参数开销保持在约 8%，改进不是依赖大幅扩容获得。")

    add_callout(
        document,
        "证据边界",
        "截图支持趋势判断，但尚不能证明统计显著性或训练公平性。正式结论仍需相同数据划分、训练设置和至少三个随机种子的受控实验。",
        accent=ORANGE,
        fill=AMBER_PALE,
    )


def add_architecture_page(document):
    add_kicker(document, "MODEL ARCHITECTURE")
    add_heading(document, "02", "CSG 模型结构")
    add_body(
        document,
        "模型接收一个五维历史场张量，以及与历史窗口对齐的 Ls 序列和无时间轴的静态地形。主干先提取历史动力特征，再由联合门控按时间、通道和空间位置进行重标定。",
    )

    table = document.add_table(rows=1, cols=3)
    for idx, text in enumerate(("变量", "张量形状", "物理含义")):
        table.rows[0].cells[idx].text = text
    rows = [
        ("x", "[B, T, C, H, W]", "历史臭氧及气象变量"),
        ("Ls", "[B, T]", "太阳黄经，单位为度"),
        ("MOLA", "[B, 1, H, W]", "静态高程，单位为米"),
        ("输出", "[B, K, 1, H, W]", "未来 K 步臭氧场"),
    ]
    for values in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(values):
            cells[idx].text = value
    style_table(table, [1600, 3200, 5064])

    add_figure(
        document,
        IMAGE_ARCHITECTURE,
        "图 1  CSG 架构示意图",
        "CSG 模型流程：输入经 ConvLSTM 和空间编码器形成历史特征，太阳黄经经 MLP 编码，MOLA 地形经卷积编码，三路信息在门控模块中融合后进入时间翻译器和空间解码器。",
        width_cm=15.5,
    )

    add_callout(
        document,
        "读图要点",
        "实现中只有一条 ConvLSTM-SimVP 主路径；右侧三路箭头表示动态特征、Ls 与 MOLA 共同决定 GATE，而不是三套独立预测网络。",
        accent=GREEN,
        fill=GREEN_PALE,
    )


def add_gate_page(document):
    add_kicker(document, "CONDITIONAL FUSION")
    add_heading(document, "03", "联合条件门控机制")
    add_body(
        document,
        "空间编码后的历史特征记为 E，形状为 [B,T,S,h,w]。门控模块将动态特征、太阳黄经和地形映射到相同的隐空间，并在非线性激活之前相加，使季节响应能够随地形和当前大气状态共同变化。",
    )

    table = document.add_table(rows=1, cols=3)
    for idx, text in enumerate(("条件分支", "编码方式", "输出作用")):
        table.rows[0].cells[idx].text = text
    rows = [
        ("动态特征 E", "1×1 卷积：S → G", "保留每个历史步的局地状态"),
        ("太阳黄经 Ls", "sin/cos 一、二阶谐波 + MLP", "连续表达季节周期，避免 359°/0° 跳变"),
        ("MOLA 地形", "高程 ÷ 10000 + 两层卷积", "提供与空间位置绑定的静态先验"),
    ]
    for values in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(values):
            cells[idx].text = value
    style_table(table, [2200, 3700, 3964])

    add_heading(document, "", "门控方程", level=2)
    add_equation_box(
        document,
        [
            "J = GELU(W_e(E) + f_Ls(Ls) + f_topo(MOLA))",
            "G = tanh(W_g(J))",
            "E_gated = E × (1 + s × G),     s ∈ [0, 1]",
        ],
    )

    add_body(
        document,
        "G 的取值限制在 [-1,1]。当 G 为正时增强对应特征，为负时抑制对应特征；可训练强度 s 经过直通式边界约束，前向始终落在 [0,1]。默认 s=0.05，因此初始尺度理论上位于 [0.95,1.05]。",
    )

    add_heading(document, "", "为什么采用近恒等初始化", level=2)
    add_bullet(document, "训练开始时保留基线主干行为，避免条件分支立即破坏已有动态表示。")
    add_bullet(document, "最终门投影采用小增益 Xavier 初始化，所有条件分支首轮反向传播即可获得有限梯度。")
    add_bullet(document, "当 s=0 且共享权重一致时，模型输出与 ConvLSTM-SimVP 基线逐元素完全相同。")

    add_callout(
        document,
        "当前限制",
        "该版本只做乘法调制。若 E 在某处接近零，Ls 与地形不能独立注入新信号；后续可在同一联合表示 J 上增加小强度加性残差分支。",
        accent=ORANGE,
        fill=AMBER_PALE,
    )


def add_overall_results_page(document):
    add_kicker(document, "EXPERIMENTAL RESULT · OVERALL")
    add_heading(document, "04", "总体 RMSE 对比")
    add_body(
        document,
        "总体指标首先回答“加入条件门是否有净收益”。当前截图中，CSG 的 RMSE 最低，ConvLSTM-SimVP 次之，SimVP 最高。下述数值为按图读数的近似值，正式报告应以平台导出的原始指标表为准。",
    )

    add_figure(
        document,
        IMAGE_OVERALL,
        "图 2  SimVP、ConvLSTM-SimVP 与 CSG 的总体 RMSE 对比",
        "横向柱状图比较三个模型总体 RMSE：CSG 约 2.12，ConvLSTM-SimVP 约 2.22，SimVP 约 2.49；RMSE 越低越好。",
        width_cm=16.5,
    )

    table = document.add_table(rows=1, cols=3)
    for idx, text in enumerate(("模型", "总体 RMSE（约）", "相对 SimVP 的变化")):
        table.rows[0].cells[idx].text = text
    rows = [
        ("CSG", "2.12", "下降约 15%"),
        ("ConvLSTM-SimVP", "2.22", "下降约 11%"),
        ("SimVP", "2.49", "参照"),
    ]
    for values in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(values):
            cells[idx].text = value
    style_table(table, [3000, 2800, 4064])

    add_heading(document, "", "结果解读", level=2)
    add_bullet(document, "ConvLSTM-SimVP 优于 SimVP，说明显式循环历史编码本身已有价值。")
    add_bullet(document, "CSG 在相同主干上进一步降低 RMSE，表明 Ls 与地形包含可利用的条件信息。")
    add_bullet(document, "CSG 相对 ConvLSTM-SimVP 的改善约为 4%–5%，属于需要多随机种子确认的中等幅度增益。")


def add_step_results_page(document):
    add_kicker(document, "EXPERIMENTAL RESULT · HORIZON")
    add_heading(document, "05", "20 步预测中的误差演化")
    add_body(
        document,
        "逐步曲线用于判断总体增益来自哪个预测区间。三种模型的误差都随预测步增长，但 CSG 在绝大多数步上保持最低，且中长预测步的优势更稳定。",
    )

    add_figure(
        document,
        IMAGE_STEPS,
        "图 3  三种模型在 20 个预测步上的 RMSE 变化",
        "折线图展示 SimVP、ConvLSTM-SimVP 和 CSG 从 Step 1 到 Step 20 的 RMSE。CSG 曲线整体最低；SimVP 在中长预测步最高。",
        width_cm=16.5,
    )

    table = document.add_table(rows=1, cols=3)
    for idx, text in enumerate(("预测区间", "曲线特征", "模型含义")):
        table.rows[0].cells[idx].text = text
    rows = [
        ("Step 1–5", "三模型同步快速上升，CSG 已建立领先", "条件先验对短期也有帮助"),
        ("Step 6–12", "CSG 与两条基线的间距持续扩大", "联合条件有助于抑制误差累积"),
        ("Step 13–20", "CSG 约趋于平台，SimVP 继续升高", "长期预测稳定性改善更明显"),
    ]
    for values in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(values):
            cells[idx].text = value
    style_table(table, [1800, 4300, 3764])

    add_callout(
        document,
        "关键观察",
        "CSG 的价值并非只集中在首个预测步；更值得关注的是中长预测步的持续分离。下一轮应报告每个 lead 的均值与标准差，而不是仅依赖单次曲线。",
        accent=GREEN,
        fill=GREEN_PALE,
    )


def add_engineering_page(document):
    add_kicker(document, "ENGINEERING REVIEW")
    add_heading(document, "06", "实现验证与风险审查")
    add_metric_strip(
        document,
        [
            ("170,338", "CSG 可训练参数", BLUE),
            ("+8.19%", "相对基线开销", ORANGE),
            ("20 / 20", "模型专用测试通过", GREEN),
        ],
    )

    add_heading(document, "", "已经验证的工程性质", level=2)
    add_bullet(document, "支持主配置 20→20、奇数空间尺寸、多输入通道及完整反向传播。")
    add_bullet(document, "Ls 周期谐波、MOLA 物理缩放、门值范围、梯度有限性均有单元测试。")
    add_bullet(document, "共享主干权重且门强度为零时，与 ConvLSTM-SimVP 基线严格等价。")
    add_bullet(document, "模型文件只依赖 torch，符合独立上传和平台 dry-run 约束。")

    add_heading(document, "", "优先处理的风险", level=2)
    table = document.add_table(rows=1, cols=3)
    for idx, text in enumerate(("问题", "可能影响", "建议")):
        table.rows[0].cells[idx].text = text
    rows = [
        (
            "地形 dtype 校验不一致",
            "float64 地形通过检查后在 float32 卷积处报错",
            "严格要求 float32，或显式转换到权重 dtype",
        ),
        (
            "纯乘法条件注入",
            "编码特征接近零时无法加入新条件信号",
            "增加小强度加性残差分支，并保留零强度兼容性",
        ),
        (
            "门控初始化过弱",
            "实测初始平均尺度变化约 0.015%，条件分支学习可能偏慢",
            "记录门统计与梯度；必要时调整强度或使用分组学习率",
        ),
        (
            "经度边界采用零填充",
            "若输入覆盖全球经度，边界可能出现人为接缝",
            "评估经度循环填充，保持纬度方向现有处理",
        ),
    ]
    for values in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(values):
            cells[idx].text = value
    style_table(table, [2450, 3500, 3914], header_fill=AMBER_PALE)

    add_callout(
        document,
        "结论尺度",
        "测试通过说明接口、形状、梯度和回归行为可靠；它不能替代科学消融，也不能证明 RMSE 改善在不同随机种子上稳定。",
        accent=ORANGE,
        fill=AMBER_PALE,
    )


def add_conclusion_page(document):
    add_kicker(document, "CONCLUSION & NEXT STEP")
    add_heading(document, "07", "结论与下一轮实验")

    add_heading(document, "", "本次汇报结论", level=2)
    add_numbered(
        document,
        "结构合理",
        "CSG 以较小参数开销在固定主干上联合利用历史状态、季节相位和静态地形，实验归因边界清晰。",
        1,
    )
    add_numbered(
        document,
        "结果积极",
        "当前截图中 CSG 的总体与逐步 RMSE 均优于两条基线，尤其在中长预测步表现更稳定。",
        2,
    )
    add_numbered(
        document,
        "结论仍需加固",
        "现有结果缺少多随机种子统计与分层分析，当前应表述为“有前景的初步增益”，而非最终科学结论。",
        3,
    )

    add_heading(document, "", "建议的三组受控实验", level=2)
    table = document.add_table(rows=1, cols=3)
    for idx, text in enumerate(("实验组", "模型变化", "回答的问题")):
        table.rows[0].cells[idx].text = text
    rows = [
        ("A · 基线", "ConvLSTM-SimVP", "固定主干可以达到什么水平？"),
        ("B · 当前 CSG", "增加 Ls + MOLA 乘法门", "条件重标定是否带来稳定增益？"),
        ("C · 残差 CSG", "在 B 上增加小强度加性条件残差", "直接注入条件特征能否进一步改善长预测步？"),
    ]
    for values in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(values):
            cells[idx].text = value
    style_table(table, [2200, 3100, 4564], header_fill=GREEN_PALE)

    add_heading(document, "", "评价协议", level=2)
    add_bullet(document, "每组至少 3 个随机种子，报告验证集 RMSE 的均值与标准差。")
    add_bullet(document, "同时报告 Step 1–20、30° Ls 分箱和 MOLA 高程四分位的 RMSE。")
    add_bullet(document, "记录 gate 的均值、标准差、最小值、最大值及强度 s，检查是否失效或饱和。")
    add_bullet(document, "保持数据划分、归一化、优化器、批量大小、训练轮数和早停策略一致。")

    add_callout(
        document,
        "下一步决策",
        "优先修复 dtype 契约，并实施加性条件残差版本；随后按 A/B/C 三组进行多种子受控实验。若增益能跨预测步、季节和地形分层保持，再进入更深入的物理解释。",
        accent=BLUE,
        fill=BLUE_PALE,
    )


def build_document():
    for path in (IMAGE_OVERALL, IMAGE_STEPS, IMAGE_ARCHITECTURE):
        if not path.is_file():
            raise FileNotFoundError(path)

    document = Document()
    configure_styles(document)
    configure_page(document)

    properties = document.core_properties
    properties.title = "CSG 条件时空门控模型组会汇报"
    properties.subject = "火星臭氧多步预测模型结构与实验结果"
    properties.author = ""
    properties.last_modified_by = ""
    properties.comments = "基于仓库实现与用户提供的三张实验截图生成。"
    properties.created = datetime(2026, 8, 26)
    properties.modified = datetime(2026, 8, 26)

    add_cover(document)
    add_page_break(document)
    add_research_page(document)
    add_page_break(document)
    add_architecture_page(document)
    add_page_break(document)
    add_gate_page(document)
    add_page_break(document)
    add_overall_results_page(document)
    add_page_break(document)
    add_step_results_page(document)
    add_page_break(document)
    add_engineering_page(document)
    add_page_break(document)
    add_conclusion_page(document)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_document()
