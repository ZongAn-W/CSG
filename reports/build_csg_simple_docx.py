import re
from datetime import datetime
from pathlib import Path

from PIL import Image
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor, Twips


SOURCE = Path(r"D:\Obsidian\仓库\我的知识图谱\项目\2026-8-26 CSG汇报\CSG组会汇报.md")
OUTPUT = SOURCE.with_name("CSG组会汇报_简洁版.docx")

FONT_LATIN = "Calibri"
FONT_CJK = "Microsoft YaHei"
COLOR_TEXT = "202124"
COLOR_MUTED = "5F6368"
COLOR_HEADING = "2E74B5"
COLOR_HEADING_DARK = "1F4D78"
COLOR_TABLE_HEADER = "F2F4F7"
COLOR_BORDER = "B7C0CE"
TABLE_WIDTH_DXA = 9360
TABLE_INDENT_DXA = 120

IMAGE_ALT = {
    "Pasted image 20260826164054.png": "CSG 模型架构：历史输入经 ConvLSTM 和空间编码器形成动态特征，太阳黄经与 MOLA 地形分别编码后共同参与门控，再进入时间翻译器和空间解码器。",
    "Pasted image 20260826164127.png": "SimVP、ConvLSTM-SimVP 与 CSG 的总体 RMSE 柱状对比，CSG 最低。",
    "Pasted image 20260826164300.png": "SimVP、ConvLSTM-SimVP 与 CSG 在 20 个预测步上的 RMSE 曲线，CSG 在多数预测步最低。",
}


def rgb(value):
    return RGBColor.from_string(value)


def set_run_font(run, size=None, color=COLOR_TEXT, bold=None, italic=None, latin=FONT_LATIN):
    run.font.name = latin
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), latin)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), latin)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT_CJK)
    if size is not None:
        run.font.size = Pt(size)
    if color:
        run.font.color.rgb = rgb(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_style_font(style, size, color=COLOR_TEXT, bold=None):
    style.font.name = FONT_LATIN
    style._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), FONT_LATIN)
    style._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), FONT_LATIN)
    style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT_CJK)
    style.font.size = Pt(size)
    style.font.color.rgb = rgb(color)
    if bold is not None:
        style.font.bold = bold


def configure_styles(document):
    normal = document.styles["Normal"]
    set_style_font(normal, 11)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    title = document.styles["Title"]
    set_style_font(title, 22, color="000000", bold=True)
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(4)
    title.paragraph_format.keep_with_next = True
    title_ppr = title._element.get_or_add_pPr()
    title_border = title_ppr.find(qn("w:pBdr"))
    if title_border is not None:
        title_ppr.remove(title_border)

    subtitle = document.styles["Subtitle"]
    set_style_font(subtitle, 11, color=COLOR_MUTED, bold=False)
    subtitle.font.italic = False
    subtitle.paragraph_format.space_before = Pt(0)
    subtitle.paragraph_format.space_after = Pt(16)
    subtitle.paragraph_format.keep_with_next = True

    heading_1 = document.styles["Heading 1"]
    set_style_font(heading_1, 16, color=COLOR_HEADING, bold=True)
    heading_1.paragraph_format.space_before = Pt(16)
    heading_1.paragraph_format.space_after = Pt(8)
    heading_1.paragraph_format.keep_with_next = True
    heading_1.paragraph_format.keep_together = True

    heading_2 = document.styles["Heading 2"]
    set_style_font(heading_2, 13, color=COLOR_HEADING, bold=True)
    heading_2.paragraph_format.space_before = Pt(12)
    heading_2.paragraph_format.space_after = Pt(6)
    heading_2.paragraph_format.keep_with_next = True
    heading_2.paragraph_format.keep_together = True

    heading_3 = document.styles["Heading 3"]
    set_style_font(heading_3, 12, color=COLOR_HEADING_DARK, bold=True)
    heading_3.paragraph_format.space_before = Pt(8)
    heading_3.paragraph_format.space_after = Pt(4)
    heading_3.paragraph_format.keep_with_next = True

    caption = document.styles["Caption"]
    set_style_font(caption, 9, color=COLOR_MUTED, bold=False)
    caption.paragraph_format.space_before = Pt(4)
    caption.paragraph_format.space_after = Pt(8)
    caption.paragraph_format.line_spacing = 1.0


def configure_page(document):
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    header = section.header
    header_p = header.paragraphs[0]
    header_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header_p.paragraph_format.space_after = Pt(0)
    set_run_font(header_p.add_run("CSG 组会汇报"), size=8.5, color=COLOR_MUTED)

    footer = section.footer
    footer_p = footer.paragraphs[0]
    footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_p.paragraph_format.space_after = Pt(0)
    page_run = footer_p.add_run()
    set_run_font(page_run, size=8.5, color=COLOR_MUTED)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for node in (begin, instruction, separate, text, end):
        page_run._r.append(node)


def append_inline(paragraph, text, size=None, color=COLOR_TEXT, bold_default=False):
    pattern = re.compile(r"(\*\*.+?\*\*|==.+?==|`.+?`)")
    cursor = 0
    for match in pattern.finditer(text):
        if match.start() > cursor:
            run = paragraph.add_run(text[cursor:match.start()])
            set_run_font(run, size=size, color=color, bold=bold_default)
        token = match.group(0)
        if token.startswith("**"):
            value = token[2:-2]
            run = paragraph.add_run(value)
            set_run_font(run, size=size, color=color, bold=True)
        elif token.startswith("=="):
            value = token[2:-2]
            run = paragraph.add_run(value)
            set_run_font(run, size=size, color=color, bold=True)
        else:
            value = token[1:-1]
            run = paragraph.add_run(value)
            set_run_font(run, size=size, color=color, bold=bold_default, latin="Consolas")
        cursor = match.end()
    if cursor < len(text):
        run = paragraph.add_run(text[cursor:])
        set_run_font(run, size=size, color=color, bold=bold_default)


def next_numbering_id(numbering, tag, attr):
    values = []
    for node in numbering.findall(qn(tag)):
        value = node.get(qn(attr))
        if value is not None:
            values.append(int(value))
    return max(values, default=0) + 1


def create_numbering(document, kind):
    numbering = document.part.numbering_part.element
    abstract_id = next_numbering_id(numbering, "w:abstractNum", "w:abstractNumId")
    num_id = next_numbering_id(numbering, "w:num", "w:numId")

    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    level = OxmlElement("w:lvl")
    level.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    num_fmt = OxmlElement("w:numFmt")
    num_fmt.set(qn("w:val"), "bullet" if kind == "bullet" else "decimal")
    level_text = OxmlElement("w:lvlText")
    level_text.set(qn("w:val"), "•" if kind == "bullet" else "%1.")
    level_jc = OxmlElement("w:lvlJc")
    level_jc.set(qn("w:val"), "left")
    ppr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "720")
    tabs.append(tab)
    indent = OxmlElement("w:ind")
    indent.set(qn("w:left"), "720")
    indent.set(qn("w:hanging"), "360")
    ppr.extend([tabs, indent])
    level.extend([start, num_fmt, level_text, level_jc, ppr])
    abstract.append(level)
    numbering.append(abstract)

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), str(abstract_id))
    num.append(abstract_ref)
    numbering.append(num)
    return num_id


def add_list_paragraph(document, text, num_id):
    paragraph = document.add_paragraph()
    ppr = paragraph._p.get_or_add_pPr()
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num_id_node = OxmlElement("w:numId")
    num_id_node.set(qn("w:val"), str(num_id))
    num_pr.extend([ilvl, num_id_node])
    ppr.append(num_pr)
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(8)
    paragraph.paragraph_format.line_spacing = 1.167
    append_inline(paragraph, text)
    return paragraph


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for key, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{key}"))
        if node is None:
            node = OxmlElement(f"w:{key}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths_dxa):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.first_child_found_in("w:tblInd")
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(TABLE_INDENT_DXA))
    tbl_ind.set(qn("w:type"), "dxa")
    layout = tbl_pr.first_child_found_in("w:tblLayout")
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(width))
        grid.append(grid_col)

    for row in table.rows:
        for index, cell in enumerate(row.cells):
            cell.width = Twips(widths_dxa[index])
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.first_child_found_in("w:tcW")
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths_dxa[index]))
            tc_w.set(qn("w:type"), "dxa")


def set_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "4")
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), COLOR_BORDER)


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), fill)


def table_widths(headers):
    first = headers[0].replace("`", "").strip()
    if first == "变量":
        return [1500, 3000, 4860]
    if first == "条件分支":
        return [1900, 3500, 3960]
    if first == "模型" and "RMSE" in headers[1]:
        return [2600, 2400, 4360]
    if first == "模型":
        return [2100, 3500, 3760]
    if first == "问题":
        return [2100, 3460, 3800]
    if first == "实验组":
        return [1800, 3200, 4360]
    count = len(headers)
    base = TABLE_WIDTH_DXA // count
    return [base] * (count - 1) + [TABLE_WIDTH_DXA - base * (count - 1)]


def add_table(document, rows):
    widths = table_widths(rows[0])
    table = document.add_table(rows=len(rows), cols=len(rows[0]))
    for row_index, values in enumerate(rows):
        row = table.rows[row_index]
        row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        if row_index == 0:
            row._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
        for column_index, value in enumerate(values):
            cell = row.cells[column_index]
            cell.text = ""
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.0
            paragraph.paragraph_format.keep_with_next = row_index < len(rows) - 1
            if row_index == 0:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                shade_cell(cell, COLOR_TABLE_HEADER)
            append_inline(paragraph, value, size=9.5, bold_default=row_index == 0)
    set_table_geometry(table, widths)
    set_table_borders(table)
    spacer = document.add_paragraph()
    spacer.paragraph_format.space_after = Pt(0)
    return table


def set_alt_text(inline_shape, title, description):
    doc_pr = inline_shape._inline.docPr
    doc_pr.set("title", title)
    doc_pr.set("descr", description)


def add_image(document, image_path):
    with Image.open(image_path) as image:
        width_px, height_px = image.size
    max_width = 6.2
    max_height = 4.45
    width = max_width
    height = width * height_px / width_px
    if height > max_height:
        height = max_height
        width = height * width_px / height_px
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.keep_with_next = True
    shape = paragraph.add_run().add_picture(str(image_path), width=Inches(width), height=Inches(height))
    set_alt_text(shape, image_path.stem, IMAGE_ALT.get(image_path.name, image_path.stem))


def add_code_block(document, lines):
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.left_indent = Inches(0.25)
    paragraph.paragraph_format.right_indent = Inches(0.25)
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(8)
    paragraph.paragraph_format.line_spacing = 1.0
    ppr = paragraph._p.get_or_add_pPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), "F7F7F7")
    ppr.append(shading)
    for index, line in enumerate(lines):
        if index:
            paragraph.add_run("\n")
        run = paragraph.add_run(line)
        set_run_font(run, size=9.5, latin="Consolas")


def split_table_row(line):
    return [cell.strip() for cell in line.strip().strip("|").split("|")]


def is_table_separator(cells):
    return all(re.fullmatch(r":?-{3,}:?", cell.replace(" ", "")) for cell in cells)


def next_content_is_image(lines, start_index):
    for line in lines[start_index:]:
        stripped = line.strip()
        if stripped:
            return re.fullmatch(r"!\[\[(.+?)\]\]", stripped) is not None
    return False


def parse_document(document, source_text):
    lines = source_text.splitlines()
    title_done = False
    subtitle_done = False
    in_code = False
    code_lines = []
    current_list_kind = None
    current_num_id = None
    index = 0

    while index < len(lines):
        stripped = lines[index].strip()

        if in_code:
            if stripped.startswith("```"):
                add_code_block(document, code_lines)
                in_code = False
                code_lines = []
            else:
                code_lines.append(lines[index])
            index += 1
            continue

        if stripped.startswith("```"):
            in_code = True
            current_list_kind = None
            index += 1
            continue

        if not stripped:
            current_list_kind = None
            index += 1
            continue

        if stripped.startswith("# ") and not title_done:
            paragraph = document.add_paragraph(style="Title")
            append_inline(paragraph, stripped[2:].strip(), size=22, color="000000", bold_default=True)
            title_done = True
            current_list_kind = None
            index += 1
            continue

        if title_done and not subtitle_done and stripped.startswith("- "):
            paragraph = document.add_paragraph(style="Subtitle")
            append_inline(paragraph, stripped[2:].strip(), size=11, color=COLOR_MUTED)
            subtitle_done = True
            current_list_kind = None
            index += 1
            continue

        if stripped.startswith("## "):
            paragraph = document.add_paragraph(style="Heading 1")
            append_inline(paragraph, stripped[3:].strip(), size=16, color=COLOR_HEADING, bold_default=True)
            current_list_kind = None
            index += 1
            continue

        if stripped.startswith("### "):
            paragraph = document.add_paragraph(style="Heading 2")
            append_inline(paragraph, stripped[4:].strip(), size=13, color=COLOR_HEADING, bold_default=True)
            current_list_kind = None
            index += 1
            continue

        if stripped.startswith("|"):
            table_rows = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                cells = split_table_row(lines[index])
                if not is_table_separator(cells):
                    table_rows.append(cells)
                index += 1
            add_table(document, table_rows)
            current_list_kind = None
            continue

        image_match = re.fullmatch(r"!\[\[(.+?)\]\]", stripped)
        if image_match:
            image_path = SOURCE.parent / image_match.group(1)
            if not image_path.is_file():
                raise FileNotFoundError(image_path)
            add_image(document, image_path)
            current_list_kind = None
            index += 1
            continue

        if stripped.startswith("**图 "):
            paragraph = document.add_paragraph(style="Caption")
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            append_inline(paragraph, stripped, size=9, color=COLOR_MUTED)
            current_list_kind = None
            index += 1
            continue

        numbered = re.match(r"^\d+\.\s+(.+)$", stripped)
        bullet = re.match(r"^-\s+(.+)$", stripped)
        if numbered or bullet:
            kind = "decimal" if numbered else "bullet"
            value = (numbered or bullet).group(1)
            if current_list_kind != kind:
                current_num_id = create_numbering(document, kind)
                current_list_kind = kind
            add_list_paragraph(document, value, current_num_id)
            index += 1
            continue

        paragraph = document.add_paragraph()
        paragraph.paragraph_format.keep_with_next = next_content_is_image(lines, index + 1)
        append_inline(paragraph, stripped)
        current_list_kind = None
        index += 1

    if in_code:
        raise ValueError("Unclosed fenced code block")


def build():
    if not SOURCE.is_file():
        raise FileNotFoundError(SOURCE)

    source_text = SOURCE.read_text(encoding="utf-8-sig")
    document = Document()
    configure_styles(document)
    configure_page(document)
    parse_document(document, source_text)

    properties = document.core_properties
    properties.title = "CSG 条件时空门控模型"
    properties.subject = "融合太阳黄经与 MOLA 地形的火星臭氧多步预测"
    properties.author = ""
    properties.last_modified_by = ""
    properties.created = datetime(2026, 8, 26)
    properties.modified = datetime(2026, 8, 26)

    document.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
